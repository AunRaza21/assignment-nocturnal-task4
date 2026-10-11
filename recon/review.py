import hashlib
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config

HEADERS = ["case_id", "grn", "outcome", "issue", "suggestion", "register", "erp", "sheet",
           "decision", "note", "decided_at", "status", "fingerprint"]
WIDTHS = [14, 8, 11, 45, 40, 40, 40, 40, 18, 25, 18, 10, 12]
YELLOW = PatternFill("solid", start_color="FFF2CC")
RED = PatternFill("solid", start_color="F8CBAD")


def sort_key(case):
    return (-case["score"], case["case_id"])


def fingerprint(case):
    text = case["issue"]
    for source in config.SOURCES:
        text += "|" + case["display"][source]
    return hashlib.sha1(text.encode()).hexdigest()[:10]


def read_old(path):
    old = {}
    if not path.exists():
        return old
    sheet = openpyxl.load_workbook(path).active
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(HEADERS, values))
        old[record["case_id"]] = record
    return old


def write_review(cases, path=config.REVIEW_FILE):
    old = read_old(path)
    todo = []
    for case in cases:
        if case["review"]:
            todo.append(case)
    todo.sort(key=sort_key)

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "review"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for case in todo:
        print_fp = fingerprint(case)
        decision = ""
        note = ""
        decided_at = ""
        status = "PENDING"
        before = old.get(case["case_id"])
        if before and before["fingerprint"] == print_fp:
            decision = before["decision"] or ""
            note = before["note"] or ""
            decided_at = before["decided_at"] or ""
            status = before["status"] or "PENDING"
        elif before and before["decision"]:
            note = f"Data changed since last review (earlier decision: {before['decision']})"
        sheet.append([
            case["case_id"], case["grn"], case["outcome"], case["issue"], case["suggestion"],
            case["display"]["register"], case["display"]["erp"], case["display"]["sheet"],
            decision, note, decided_at, status, print_fp,
        ])
        color_row(sheet, sheet.max_row, case)

    format_sheet(sheet, len(todo))
    config.OUTPUT.mkdir(exist_ok=True)
    book.save(path)
    return len(todo)


def color_row(sheet, row_number, case):
    for source in config.SOURCES:
        if source not in case["highlight"]:
            continue
        column = HEADERS.index(source) + 1
        cell = sheet.cell(row=row_number, column=column)
        if cell.value == "MISSING":
            cell.fill = RED
        else:
            cell.fill = YELLOW


def format_sheet(sheet, count):
    for number, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["M"].hidden = True
    sheet.freeze_panes = "A2"
    options = '"' + ",".join(config.DECISIONS) + '"'
    rule = DataValidation(type="list", formula1=options, allow_blank=True)
    sheet.add_data_validation(rule)
    column = get_column_letter(HEADERS.index("decision") + 1)
    rule.add(f"{column}2:{column}{count + 1}")


def apply_decisions(path=config.REVIEW_FILE):
    if not path.exists():
        return 0, []
    book = openpyxl.load_workbook(path)
    sheet = book.active
    applied = 0
    errors = []
    for row in sheet.iter_rows(min_row=2):
        cells = dict(zip(HEADERS, row))
        decision = str(cells["decision"].value or "").strip().upper()
        if decision == "" or cells["status"].value == "RESOLVED":
            continue
        if decision not in config.DECISIONS:
            errors.append(f"{cells['case_id'].value}: '{decision}' is not one of {', '.join(config.DECISIONS)}")
            continue
        cells["decision"].value = decision
        cells["decided_at"].value = datetime.now().strftime("%Y-%m-%d %H:%M")
        cells["status"].value = "RESOLVED"
        applied += 1
    book.save(path)
    return applied, errors


def read_pending(path=config.REVIEW_FILE):
    pending = []
    if not path.exists():
        return pending
    sheet = openpyxl.load_workbook(path).active
    for values in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(HEADERS, values))
        if record["status"] != "RESOLVED":
            pending.append(record)
    return pending
